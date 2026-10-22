from openpyxl import load_workbook
from domain.Estudiante import Estudiante


def encontrar_fila_encabezado(hoja): #funcion para encontrar en que fila se encuentran los encabezados
    for numero_fila, filas in enumerate(hoja.iter_rows(), start=1):
        for celda in filas:
            if isinstance(celda.value, str):
                texto = celda.value.strip().upper()

                if "DOCUMENTO" in texto:
                    return numero_fila
                

def obtener_columnas(hoja, fila_encabezado): #funcion para obtener diccionario de donde estan los encabezados y sus posiciones
    columnas = {}

    for indice, celda in enumerate(hoja[fila_encabezado]):
        if isinstance(celda.value, str):
            texto = celda.value.strip().upper()

            if "DOCUMENTO" in texto:    
                columnas["documento"] = indice

            if "NOMBRE" in texto:
                columnas["nombre"] = indice

            if "COLEGIO" in texto:
                columnas["colegio"] = indice
                
            if "GRUPO" in texto:
                columnas["grupo"] = indice
            
            if "CURSO" in texto:
                columnas["grado"] = indice
                
            if "MUNICIPIO" in texto:
                columnas["ciudad"] = indice
    return columnas




def agregar_estudiantes(ruta_archivo, nombre_hoja): #funcion para agregar estudiantes 
    workbook = load_workbook(ruta_archivo)
    hoja = workbook[nombre_hoja]
    
    fila_encabezado = encontrar_fila_encabezado(hoja)
    columnas = obtener_columnas(hoja, fila_encabezado)
    estudiantes = []
    
    for fila in hoja.iter_rows(min_row=fila_encabezado + 1):
        
        
        documento = fila[columnas["documento"]].value
        nombre = fila[columnas["nombre"]].value
        colegio = fila[columnas["colegio"]].value
        grupo = fila[columnas["grupo"]].value
        grado = fila[columnas["grado"]].value
        ciudad = fila[columnas["ciudad"]].value
        
        if documento is None or nombre is None:
            continue
        
        documento = str(documento)
        nombre = str(nombre)
        colegio = str(colegio)
        grupo = str(grupo)
        grado = grado.replace("°", "").split()[-1]
        ciudad = str(ciudad)
        
        
        estudiante = Estudiante(
            documento,
            nombre,
            colegio,
            grupo,
            grado,
            ciudad
        )
        
        estudiantes.append(estudiante)
    return estudiantes